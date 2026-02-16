from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile


def _safe_read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception:  # noqa: BLE001
        return "[未安装 pypdf，无法解析 PDF 文本]"

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except Exception:  # noqa: BLE001
        return "[未安装 python-docx，无法解析 DOCX 文本]"

    document = docx.Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _read_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except Exception:  # noqa: BLE001
        return "[未安装 python-pptx，无法解析 PPTX 文本]"

    prs = Presentation(str(path))
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                lines.append(text)
    return "\n".join(lines)


def _read_xlsx(path: Path) -> str:
    try:
        import pandas as pd
    except Exception:  # noqa: BLE001
        return "[未安装 pandas/openpyxl，无法解析 XLSX 文本]"

    lines: list[str] = []
    sheets = pd.read_excel(path, sheet_name=None)
    for sheet_name, frame in sheets.items():
        lines.append(f"## Sheet: {sheet_name}")
        lines.append(frame.head(200).to_csv(index=False))
    return "\n".join(lines)


def _read_image(path: Path) -> str:
    try:
        import pytesseract
        from PIL import Image
    except Exception:  # noqa: BLE001
        return f"[图片文件 {path.name}：未安装 OCR 依赖，跳过文本提取]"

    image = Image.open(path)
    text = pytesseract.image_to_string(image, lang="chi_sim+eng")
    return text or f"[图片文件 {path.name}：未提取到可见文本]"


def _read_doc_legacy(path: Path) -> str:
    # .doc is a binary format. Use antiword if available.
    import shutil
    import subprocess

    antiword = shutil.which("antiword")
    if not antiword:
        return "[DOC 旧格式解析需要 antiword，当前环境不可用]"

    try:
        output = subprocess.check_output([antiword, str(path)], text=True)
    except Exception as exc:  # noqa: BLE001
        return f"[DOC 旧格式解析失败: {exc}]"
    return output


def _read_zip(path: Path) -> str:
    chunks: list[str] = []
    with ZipFile(path, "r") as zf:
        for name in zf.namelist()[:50]:
            if name.lower().endswith(".txt"):
                chunks.append(f"# {name}")
                chunks.append(zf.read(name).decode("utf-8", errors="ignore"))
    return "\n".join(chunks) if chunks else "[ZIP 中未发现可直接解析文本文件]"


def load_document(path: str) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".xml", ".html"}:
        return _safe_read_text(p)
    if suffix == ".pdf":
        return _read_pdf(p)
    if suffix == ".docx":
        return _read_docx(p)
    if suffix == ".doc":
        return _read_doc_legacy(p)
    if suffix in {".pptx", ".ppt"}:
        return _read_pptx(p)
    if suffix in {".xlsx", ".xls"}:
        return _read_xlsx(p)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        return _read_image(p)
    if suffix == ".zip":
        return _read_zip(p)
    return f"[暂不支持的文件类型: {p.name}]"


def combine_documents(paths: list[str], limit_chars: int = 40_000) -> str:
    if not paths:
        return ""

    chunks: list[str] = []
    for raw in paths:
        path = Path(raw)
        if not path.exists():
            continue
        text = load_document(str(path))
        chunks.append(f"\n\n# 文件: {path.name}\n{text[:8000]}")

    merged = "\n".join(chunks)
    return merged[:limit_chars]
